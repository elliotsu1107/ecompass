"""Read CSV and Excel reports with header detection."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any

@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, Any]]
    encoding: str | None = None
    header_row: int = 0
    sheet: str | None = None


def _clean(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _header_score(headers: list[str]) -> int:
    terms = ("日期", "主体ID", "商品ID", "支付金额", "花费", "总成交金额", "店铺")
    return sum(any(term in h for h in headers) for term in terms)


def _from_matrix(matrix: list[list[Any]], encoding: str | None = None, sheet: str | None = None) -> Table:
    best = max(range(min(20, len(matrix))), key=lambda i: _header_score([_clean(x) for x in matrix[i]]), default=0)
    headers = [_clean(x) or f"未命名列{i + 1}" for i, x in enumerate(matrix[best])]
    rows = []
    for values in matrix[best + 1:]:
        values = list(values)
        if not any(_clean(x) for x in values):
            continue
        values += [None] * (len(headers) - len(values))
        rows.append({h: values[i] for i, h in enumerate(headers)})
    return Table(headers, rows, encoding, best, sheet)


def read_table(path: str | Path) -> Table:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030", "gbk"):
            try:
                text = raw.decode(encoding)
                matrix = list(csv.reader(text.splitlines()))
                return _from_matrix(matrix, encoding)
            except UnicodeDecodeError:
                continue
        raise UnicodeDecodeError("unknown", raw, 0, 1, "unsupported report encoding")
    if suffix in (".xlsx", ".xls"):
        matrix = None
        sheet_name = None
        try:
            import openpyxl
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_name = book.sheetnames[0]
            try:
                matrix = [list(row) for row in book[sheet_name].iter_rows(values_only=True)]
            finally:
                book.close()
        except Exception:
            try:
                import pandas as pd
                book = pd.ExcelFile(path)
                try:
                    sheet_name = book.sheet_names[0]
                    matrix = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=20).values.tolist()
                    matrix += pd.read_excel(path, sheet_name=sheet_name, header=None, skiprows=20).values.tolist()
                finally:
                    book.close()
            except Exception as exc:
                raise RuntimeError("读取 Excel 需要安装 openpyxl 或 pandas/xlrd") from exc
        return _from_matrix(matrix or [], sheet=sheet_name)
    raise ValueError(f"不支持的报表格式: {suffix}")
